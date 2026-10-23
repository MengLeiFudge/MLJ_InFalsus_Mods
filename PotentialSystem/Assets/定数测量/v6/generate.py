"""生成 v6 冻结模型和当前曲库表，保留旧 v4/v5 及训练缓存。

在仓库根目录运行 python3 PotentialSystem/Assets/定数测量/v6/generate.py。
拟合与读谱复用 .codex/chart-constant-estimation 中的已有研究资料和依赖缓存；
v6_model.py 对冻结模型的推理只依赖标准库。所有输出已存在时拒绝覆盖。
"""
from __future__ import annotations

import csv
from datetime import datetime, timezone
from hashlib import sha256
import json
import os
from pathlib import Path
import sys

os.environ.setdefault('OPENBLAS_NUM_THREADS', '1')
HERE = Path(__file__).resolve().parent
ROOT = HERE.parents[3]
RESEARCH = ROOT / '.codex/chart-constant-estimation'
OPS = RESEARCH / 'ops-ranking'
sys.path[:0] = [str(OPS), str(RESEARCH)]
import analyze

analyze.use_cached_packages()
import numpy as np
import per_difficulty as pd
import ml_newfeat as mn
import ml_small_nonlinear as sn
import ranking_features as rf
import sim
import chart_rank
import hand
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter
from v6_model import allowed_constants, predict

DIFFICULTIES = {1: 'MIN', 2: 'EVO', 4: 'ULT', 8: 'FBD'}
SOURCE_NAMES = {'MIN': 'minimal', 'EVO': 'evolved', 'ULT': 'ultimate', 'FBD': 'forbidden'}
SUFFIXES = {'MIN': '0', 'EVO': '1', 'ULT': '2', 'FBD': '3'}


def read_json(path):
    return json.loads(path.read_text(encoding='utf-8'))


def digest(path):
    return sha256(path.read_bytes()).hexdigest()


def write_json(path, value):
    path.write_text(json.dumps(value, ensure_ascii=False, indent=2, allow_nan=False) + '\n', encoding='utf-8')


def fit_structure(rows, columns, difficulty):
    """MIN/EVO 保留自身结构等级单位；ULT/FBD 仅由本难度参照标定缩放。"""
    fitted = pd.Model(columns).fit(rows, beta=1.0 if difficulty in ('MIN', 'EVO') else None)
    return {'columns': columns, 'weights': fitted.b.tolist(), 'intercept': float(fitted.a),
            'beta': float(fitted.beta), 'centers': {str(k): v for k, v in fitted.center.items()},
            'means': {str(k): v for k, v in fitted.mu.items()}}


def fit_ensemble(rows):
    """固定 v5 的七项特征与正则强度，单独拟合一个难度，移除难度指示。"""
    columns = list(sn.COLS)
    x = np.log1p(np.array([[r['f'][c] for c in columns] for r in rows], float))
    levels = np.array([r['official_level'] for r in rows], float)
    y = np.array([r['community'] for r in rows], float) - levels
    mean, scale = x.mean(0), x.std(0) + 1e-9
    design = np.column_stack([np.ones(len(rows)), levels, (x - mean) / scale])
    regularization = np.eye(design.shape[1])
    regularization[0, 0] = 0
    weights = np.linalg.solve(design.T @ design + regularization, design.T @ y)
    raw = np.column_stack([x, levels])
    kmean, kscale = raw.mean(0), raw.std(0) + 1e-9
    centers = np.clip((raw - kmean) / kscale, -5, 5)
    distance = ((centers[:, None, :] - centers[None, :, :]) ** 2).sum(2)
    kernel = np.exp(-.3 * distance / centers.shape[1])
    alpha = np.linalg.solve(kernel + .3 * np.eye(len(rows)), y - y.mean())
    return {'columns': columns, 'minimum_level': int(levels.min()), 'maximum_level': int(levels.max()),
            'training_count': len(rows),
            'linear': {'mean': mean.tolist(), 'scale': scale.tolist(), 'weights': weights.tolist()},
            'kernel': {'gamma': .3, 'mean': kmean.tolist(), 'scale': kscale.tolist(),
                       'centers': centers.tolist(), 'weights': alpha.tolist(), 'target_mean': float(y.mean())}}


def fit_candidate(rows, difficulty, columns, kind):
    if any(row['difficulty'] != difficulty for row in rows):
        raise ValueError('不能把其他难度的数据传入独立模型')
    model = {'difficulty': difficulty, 'kind': kind, 'training_ids': [r['chart_id'] for r in rows],
             'structure': fit_structure(rows, columns, difficulty)}
    if kind == 'ensemble':
        model['ensemble'] = fit_ensemble([r for r in rows if r['refined']])
    return model


def statistics(rows, values, reference='community'):
    """仅比较同难度同等级、参照不并列的谱面对；预测并列计半。"""
    error = [abs(values[r['chart_id']] - r[reference]) for r in rows]
    pairs = agree = tied = reverse = 0
    for index, left in enumerate(rows):
        for right in rows[index + 1:]:
            if (left['difficulty'], left['official_level']) != (right['difficulty'], right['official_level']):
                continue
            delta = left[reference] - right[reference]
            if abs(delta) < .05:
                continue
            pairs += 1
            predicted = values[left['chart_id']] - values[right['chart_id']]
            if abs(predicted) < 1e-9:
                tied += 1
            elif delta * predicted > 0:
                agree += 1
            else:
                reverse += 1
    return {'charts': len(rows), 'mae': sum(error) / len(error) if error else None,
            'within_0_1': sum(e <= .100000001 for e in error),
            'within_0_2': sum(e <= .200000001 for e in error),
            'over_0_3': sum(e > .300000001 for e in error),
            'comparable_pairs': pairs, 'same_direction': agree, 'ties': tied, 'opposite': reverse,
            'pair_agreement': (agree + .5 * tied) / pairs if pairs else None}


def choose_models(rows, selections):
    models, evaluation, leave_one = {}, {}, {}
    for difficulty in DIFFICULTIES.values():
        group = [r for r in rows if r['difficulty'] == difficulty]
        columns = selections[difficulty]
        if difficulty in ('MIN', 'EVO'):
            models[difficulty] = fit_candidate(group, difficulty, columns, 'structure')
            evaluation[difficulty] = {'selected': 'structure', 'reference_count': 0,
                                      'limitation': '无细分参照；结构估计，不报告小数精度'}
            continue
        reference = [r for r in group if r['refined']]
        candidates, predictions = {}, {}
        for kind in ('structure', 'ensemble'):
            values = {}
            for song in sorted({r['song_id'] for r in reference}):
                training = [r for r in group if r['song_id'] != song]
                model = fit_candidate(training, difficulty, columns, kind)
                for row in reference:
                    if row['song_id'] == song:
                        values[row['chart_id']] = predict(model, row)[1]
            candidates[kind] = statistics(reference, values)
            predictions[kind] = values
        # 用户要求趋势大致一致，优先同级趋势；同趋势时再看绝对误差。
        selected = max(candidates, key=lambda key: (candidates[key]['pair_agreement'], -candidates[key]['mae']))
        models[difficulty] = fit_candidate(group, difficulty, columns, selected)
        leave_one.update(predictions[selected])
        evaluation[difficulty] = {'selected': selected, 'reference_count': len(reference), 'candidates': candidates}
        print('模型', difficulty, selected, candidates, flush=True)
    return models, evaluation, leave_one


def extract_features(notes, difficulty, required):
    """新谱面复用已有读谱口径；只导出本难度模型需要的特征。"""
    start = min(n['time'] for n in notes) / 1000
    duration = max(.001, max(n['end'] for n in notes) / 1000 - start)
    costs = hand.solve(chart_rank.keyboard_events(notes), beam=32) if any(c.startswith('discount') for c in required) else []
    values = rf.keyboard_features(notes, start, duration, costs)
    values.update(rf.action_features(notes, start, duration))
    values.update(rf.sky_features(notes, start, duration))
    values.update(rf.temporal_sky_features(notes, start, duration))
    channels, _ = sim.keyboard_channels(notes, difficulty)
    channels['keyall'] = channels['key'] + channels['lock'] + channels['rhythm'] + channels['wide']
    if 'sim_air_h0.5_top' in required:
        mouse, _ = sim.mouse_channel(notes)
        channels['air'] = mouse + sim.switch_channel(notes)
    for column in required:
        if column.startswith('sim_'):
            channel, half_life, stat = column[4:].split('_')
            values[column] = sim.aggregate(channels[channel], float(half_life[1:]))[stat]
    values.update(mn.feats(notes))
    return {column: float(values[column]) for column in sorted(required)}


def current_rows(training, models):
    data, mappings, bundles = analyze.load_inputs()
    reader = sys.modules['game_catalog_reader']
    aa = analyze.GAME / 'infalsus_Data/StreamingAssets/aa'
    catalog = reader.BinaryCatalog(aa / 'catalog.bin')
    resolver = reader.UnityBundleResolver(aa / 'StandaloneWindows64', catalog)
    location = reader._find_asset_location(catalog, 'release/DynamicStringMapping.asset', 'DynamicStringMapping')
    obj, _ = resolver.resolve(location)
    titles = obj.read_typetree()['songIdTitleTypeMapping']
    old = {r['chart_id']: r for r in training}
    source = read_json(ROOT / 'PotentialSystem/Assets/AnXiaoLiConstants.json')
    output, ids = [], set()
    for song in data['allSongInfo']:
        for chart in song['ChartInfos']:
            if not chart['Available']:
                continue
            cid = chart['Id']
            if not cid or cid in ids:
                raise ValueError('当前曲库谱面身份为空或重复')
            ids.add(cid)
            difficulty = DIFFICULTIES[chart['Difficulty']]
            model = models[difficulty]
            required = set(model['structure']['columns'])
            if 'ensemble' in model:
                required.update(model['ensemble']['columns'])
            path, expected_length = mappings[cid + '.spc']
            raw = path.read_bytes()
            if len(raw) != expected_length:
                raise ValueError(f'{cid} 谱面长度与映射不符')
            notes, content_hash = analyze.decode(raw, cid + '.spc')
            previous = old.get(cid)
            reuse = previous is not None and previous['content_sha256'] == content_hash
            f = {c: previous['f'][c] for c in sorted(required)} if reuse else extract_features(notes, difficulty, required)
            sid, level = song['Id']['Value'], chart['Rating']
            community = source.get(str(sid), {}).get(SOURCE_NAMES[difficulty])
            row = {'song_id': sid, 'chart_id': cid, 'base_name': song['BaseName'],
                   'title': reader._mapping_lookup(titles, sid, 'songIdTitleTypeMapping') or song['BaseName'],
                   'difficulty': difficulty, 'official_level': level, 'community': community,
                   'new_chart': previous is None, 'features_reused': reuse, 'f': f,
                   'sam_sha256': sha256(raw).hexdigest(), 'content_sha256': content_hash,
                   'training_community': previous['community'] if previous and previous['refined'] else None,
                   'training_level': previous['official_level'] if previous else None}
            estimate, constant, method = predict(model, row)
            row.update(estimate=estimate, constant=constant, method=method)
            row['reference_in_range'] = community is not None and level <= community <= level + .9 + 1e-9
            row['status'] = '新谱面；' if row['new_chart'] else ''
            if method == '结构估计':
                row['status'] += '结构估计，缺少本级细分依据'
            else:
                row['status'] += '社区回归'
            if community is not None and not row['reference_in_range']:
                row['status'] += '；社区值超出当前官方等级'
            if row['new_chart'] and community == level:
                row['status'] += '；社区仍为整数等级，细分未确认'
            output.append(row)
            if not reuse or len(output) % 50 == 0:
                print('当前曲库', len(output), cid, '缓存' if reuse else '重新读谱', constant, flush=True)
    paths = [aa / 'catalog.bin'] + [aa / 'StandaloneWindows64' / name
             for name in sorted(set(bundles) | resolver.used_bundles)]
    return output, {'game_commit': data.get('CommitId'), 'metadata_sha256': {str(p): digest(p) for p in paths}}


def old_workbook(path):
    book = load_workbook(path, read_only=True, data_only=True)
    try:
        return {r[1] + SUFFIXES[r[3]]: r[6] for r in list(book['定数'].values)[1:]}
    finally:
        book.close()


def export_tables(rows, metrics, comparison):
    headers = ['ID', 'Base Name', 'Title', 'Difficult', 'Level', '安小璃当前定数', 'v6定数',
               'v6减社区', '连续估计', '留一定数', '训练社区定数', 'v4定数', 'v5定数', '状态', '谱面ID', '训练官方等级']
    def cells(row):
        community = row['community']
        return [row['song_id'], row['base_name'], row['title'], row['difficulty'], row['official_level'],
                community, row['constant'], None if community is None else round(row['constant'] - community, 1),
                row['estimate'], row.get('leave_one'), row['training_community'], row.get('v4'), row.get('v5'),
                row['status'], row['chart_id'], row['training_level']]
    ordered = sorted(rows, key=lambda r: (-r['constant'], r['difficulty'], r['chart_id']))
    book = Workbook()
    book.remove(book.active)
    for title, group in [('定数', ordered), ('新增谱面', [r for r in ordered if r['new_chart']]),
                         ('偏差与参照冲突', [r for r in ordered if r['community'] is not None
                                      and (abs(r['constant'] - r['community']) > .300000001 or not r['reference_in_range'])])]:
        sheet = book.create_sheet(title)
        sheet.append(headers)
        for row in group:
            sheet.append(cells(row))
        for index, width in enumerate((8, 25, 40, 12, 10, 18, 12, 14, 16, 14, 18, 12, 12, 65, 28, 16), 1):
            sheet.column_dimensions[get_column_letter(index)].width = width
        for line in sheet.iter_rows(min_row=2):
            for index in range(5, 13):
                line[index].number_format = '0.0000' if index == 8 else '0.0'
        sheet.freeze_panes = 'A2'
        sheet.auto_filter.ref = sheet.dimensions
    sheet = book.create_sheet('趋势与误差')
    sheet.append(['范围', '谱面数', '平均绝对差', '差≤0.1', '差≤0.2', '差>0.3', '可比对数', '同向', '并列', '反向', '同级一致率'])
    for name, value in metrics.items():
        sheet.append([name, *[value[key] for key in ('charts', 'mae', 'within_0_1', 'within_0_2', 'over_0_3',
                      'comparable_pairs', 'same_direction', 'ties', 'opposite', 'pair_agreement')]])
    sheet.column_dimensions['A'].width = 45
    for index in range(2, 12):
        sheet.column_dimensions[get_column_letter(index)].width = 17
    sheet = book.create_sheet('反向谱面对')
    sheet.append(['难度', '官方等级', '社区较难谱面', '社区定数', 'v6定数', '社区较易谱面', '社区定数', 'v6定数'])
    reverse = []
    eligible = [r for r in rows if r['reference_in_range']]
    for index, left in enumerate(eligible):
        for right in eligible[index + 1:]:
            if (left['difficulty'], left['official_level']) != (right['difficulty'], right['official_level']):
                continue
            high, low = sorted((left, right), key=lambda r: r['community'], reverse=True)
            if high['community'] - low['community'] > .05 and high['constant'] < low['constant'] - 1e-9:
                reverse.append((high['community'] - low['community'], high, low))
    for _, high, low in sorted(reverse, key=lambda item: -item[0]):
        sheet.append([high['difficulty'], high['official_level'], high['title'], high['community'], high['constant'],
                      low['title'], low['community'], low['constant']])
    for index in range(1, 9):
        sheet.column_dimensions[get_column_letter(index)].width = 38 if index in (3, 6) else 15
    sheet = book.create_sheet('说明')
    notes = [
        'v6：四个难度独立拟合；最终值严格属于官方等级 L 的合法分档。',
        'L<7：L/L+0.5；L=7或8：L/L+0.5/L+0.8；L>=9：L到L+0.9，步长0.1。最近分档，等距取高，不跨级。',
        'MIN/EVO：各自四项特征的非负结构回归，按本难度同级居中，beta=1；未借用ULT/FBD参数。没有细分参照，不能宣称小数准确。',
        'ULT/FBD：分别比较既有结构公式和固定七特征的线性/核岭均值，按歌曲留一的同级趋势选取，绝对误差只用于同趋势时比较。',
        '社区回归训练等级之外使用本难度结构公式，缺少同级中心时取 L+0.45，再按合法档位输出。',
        '训练冻结旧300张资料，其中116张有社区细分；不用feedback修订。新曲不参与训练，当前安小璃表只作对照。',
        '留一定数对应旧训练社区列，不能直接当作针对更新后社区值的验证；候选选择与既有特征选择仍有乐观偏差。',
        '误差与趋势均按最终分档值统计；v5旧报告使用未分档连续值，不能混用其误差与离群数量。',
        '趋势只比较同难度、同官方等级内社区不并列的谱面对；预测并列计半。无可比对数表示社区缺少区分度。',
        '硬边界不能证明官方整数等级正确；社区超出当前官方等级的条目单列，不强行拟合。',
        '当前游戏内萌泪酱仍用已发布的v4表；本工作簿和model.json尚未嵌入DLL。',
        '不保证低档数量均匀；粗分档会增加并列，且在本级顶部产生向下量化误差。',
        'v4/v5对照为原工作簿值，官方等级或社区版本有变化时不宜直接比较准确率。',
        'v4/v5差异汇总：' + json.dumps(comparison, ensure_ascii=False)]
    for note in notes:
        sheet.append([note])
    sheet.column_dimensions['A'].width = 145
    for sheet in book:
        for cell in sheet[1]:
            cell.font = Font(bold=True)
        sheet.freeze_panes = 'A2'
    book.save(HERE.parent / 'chart-constants-v6.xlsx')
    with (HERE / 'chart-constants-v6.csv').open('w', newline='', encoding='utf-8-sig') as stream:
        writer = csv.writer(stream)
        writer.writerow(headers)
        writer.writerows(cells(row) for row in ordered)


def main():
    outputs = [HERE / name for name in ('model.json', 'results.json', 'training-data.json', 'chart-constants-v6.csv')]
    outputs.append(HERE.parent / 'chart-constants-v6.xlsx')
    if any(path.exists() for path in outputs):
        raise FileExistsError('v6 输出已存在；请先归档，避免覆盖已有结果')
    required = ['formula-features.json', 'sim-features.json', 'wide-features.json', 'notes-cache.json', 'per-difficulty-selection.json']
    for name in required:
        if not (OPS / name).is_file():
            raise FileNotFoundError(f'缺少既有研究输入：{OPS / name}')
    inputs = [OPS / name for name in required] + list(OPS.glob('*.py'))
    inputs += [RESEARCH / 'analyze.py', RESEARCH / 'v3/engine.py', RESEARCH / 'extract-workload.py',
               RESEARCH / 'tool-comparison/comparison.json', HERE / 'generate.py', HERE / 'v6_model.py',
               ROOT / 'PotentialSystem/Assets/AnXiaoLiConstants.json',
               HERE.parent / 'chart-constants-v4.xlsx', HERE.parent / 'chart-constants-v5.xlsx', OPS / 'ml-final.json']
    hashes = {str(path.relative_to(ROOT)): digest(path) for path in inputs}
    training = mn.rows_all()
    selections = read_json(OPS / 'per-difficulty-selection.json')
    models, evaluation, leave_one = choose_models(training, selections)
    rows, game = current_rows(training, models)
    v4, v5 = (old_workbook(HERE.parent / f'chart-constants-v{version}.xlsx') for version in (4, 5))
    for row in rows:
        row.update(v4=v4.get(row['chart_id']), v5=v5.get(row['chart_id']), leave_one=leave_one.get(row['chart_id']))
        if row['constant'] not in allowed_constants(row['official_level']):
            raise ValueError(f"{row['chart_id']} 不符合分档")
    references = [r for r in training if r['refined']]
    direct = {r['chart_id']: predict(models[r['difficulty']], r)[1] for r in references}
    metrics = {'旧参照：v6直接预测': statistics(references, direct),
               '旧参照：v6按歌曲留一': statistics(references, leave_one)}
    old_v5 = read_json(OPS / 'ml-final.json')
    metrics['旧参照：v5保存的按歌曲留一'] = statistics(references, {key: round(value, 1) for key, value in old_v5['loo'].items()})
    current_values = {r['chart_id']: r['constant'] for r in rows}
    for difficulty in DIFFICULTIES.values():
        group = [r for r in rows if r['difficulty'] == difficulty and r['reference_in_range']]
        metrics['当前安小璃：' + difficulty] = statistics(group, current_values)
    comparison = {}
    for version, previous in (('v4', v4), ('v5', v5)):
        differences = [abs(row['constant'] - previous[row['chart_id']]) for row in rows if row['chart_id'] in previous]
        comparison[version] = {'common': len(differences), 'changed': sum(d > 1e-9 for d in differences),
                               'mean_absolute_difference': sum(differences) / len(differences), 'maximum_difference': max(differences)}
    for name, expected in hashes.items():
        if digest(ROOT / name) != expected:
            raise ValueError(f'计算期间输入发生变化：{name}')
    model_document = {'schema_version': 1, 'version': 'menglei-v6', 'quantization': 'nearest-within-official-level-ties-up', 'models': models}
    training_columns = set(sn.COLS).union(*(set(columns) for columns in selections.values()))
    frozen_training = [{**{key: row[key] for key in ('chart_id', 'song_id', 'title', 'difficulty', 'official_level', 'community', 'refined', 'content_sha256')},
                        'f': {c: row['f'][c] for c in sorted(training_columns)}} for row in training]
    results = {'generated_at': datetime.now(timezone.utc).isoformat(), 'input_sha256': hashes, 'game': game,
               'evaluation': evaluation, 'metrics': metrics, 'comparison': comparison, 'charts': rows}
    write_json(HERE / 'model.json', model_document)
    write_json(HERE / 'training-data.json', frozen_training)
    write_json(HERE / 'results.json', results)
    export_tables(rows, metrics, comparison)
    print('RESULT', HERE.parent / 'chart-constants-v6.xlsx', flush=True)
    print('SUMMARY', json.dumps({'count': len(rows), 'new': sum(r['new_chart'] for r in rows), 'metrics': metrics,
                                'comparison': comparison, 'model_bytes': (HERE / 'model.json').stat().st_size}, ensure_ascii=False), flush=True)


if __name__ == '__main__':
    main()
